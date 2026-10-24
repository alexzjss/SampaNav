"""Constrói o modelo de risco por célula H3 a partir dos dados criminais da SSP-SP.

Etapas: (1) extrai ocorrências de São Paulo das planilhas (cache parquet);
(2) classifica em categorias; (3) filtra/deduplica; (4) decaimento temporal;
(5) agrega em H3 res 9 com suavização k=1; (6) normaliza em escala log (ver normalize_risk).

Uso: python -m pipeline.build_risk [--force]
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import unicodedata
from collections import Counter

import numpy as np
import pandas as pd

from pipeline import config as C

OCC_PATH = C.DATA_PROCESSED / "occurrences.parquet"
RISK_PATH = C.DATA_PROCESSED / "risk_hex.parquet"
META_PATH = C.DATA_PROCESSED / "risk_meta.json"


# --------------------------------------------------------------------------
# Funções puras
# --------------------------------------------------------------------------
def norm_text(s) -> str:
    """Minúsculas, sem acentos, espaços colapsados. None/NaN -> ''."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


def _has(text: str, needles) -> bool:
    return any(n in text for n in needles)


def classify(rubrica, conduta, natureza, tipolocal, subtipolocal="") -> str | None:
    """Retorna a categoria do crime ou None se deve ser descartado."""
    rub, con, nat, loc, sub = (norm_text(x) for x in (rubrica, conduta, natureza, tipolocal, subtipolocal))
    if loc not in C.TIPOLOCAL_ALLOW or _has(sub, C.SUBTIPOLOCAL_DENY):
        return None
    if _has(rub, C.RUBRICA_LETAL):
        return "letal"
    if _has(rub, C.RUBRICA_HOMICIDIO):
        if _has(nat, C.NATUREZA_LETAL):
            return "letal"
        return "lesao" if _has(nat, C.NATUREZA_LESAO) else None
    if _has(rub, C.RUBRICA_LESAO):
        return "lesao" if _has(nat, C.NATUREZA_LESAO) else None
    is_veh = _has(con, C.CONDUTA_VEICULO) or _has(nat, C.NATUREZA_VEICULO)
    if _has(rub, C.RUBRICA_ROUBO):
        if _has(nat, C.NATUREZA_LETAL) or _has(con, C.CONDUTA_LETAL_ROUBO):
            return "letal"  # latrocínio
        if _has(con, C.CONDUTA_CARGA) or _has(nat, C.NATUREZA_CARGA):
            return None
        if _has(con, C.CONDUTA_RESIDENCIA) and not is_veh:
            return None
        return "roubo_veiculo" if is_veh else "roubo_pessoa"
    if _has(rub, C.RUBRICA_FURTO):
        if _has(con, C.CONDUTA_CARGA) or _has(nat, C.NATUREZA_CARGA):
            return None
        if _has(con, C.CONDUTA_RESIDENCIA) and not is_veh:
            return None
        return "furto_veiculo" if is_veh else "furto_pessoa"
    return None


def period_from_hour(hour: int | None) -> str | None:
    """Período do dia a partir da hora (0-23); None se desconhecida."""
    if hour is None or not (0 <= hour <= 23):
        return None
    for name, (a, b) in C.PERIOD_HOURS.items():
        if a <= hour < b:
            return name
    return None


def period_from_desc(desc) -> str | None:
    """Fallback usando DESC_PERIODO ('De madrugada', 'Pela manhã', 'A tarde', 'A noite')."""
    d = norm_text(desc)
    for key, name in C.DESC_PERIODO_MAP.items():
        if key in d:
            return name
    return None


def parse_hour(h) -> int | None:
    """Extrai a hora de datetime.time, 'HH:MM[:SS]' ou 'NULL'."""
    if isinstance(h, (dt.time, dt.datetime)):
        return h.hour
    if isinstance(h, str):
        parts = h.strip().split(":")
        if parts and parts[0].isdigit():
            v = int(parts[0])
            return v if 0 <= v <= 23 else None
    return None


def decay_weight(age_days, half_life_days: float = C.HALF_LIFE_DAYS):
    """w_t = 0.5 ** (idade / meia-vida). Aceita escalar ou array."""
    return 0.5 ** (np.clip(age_days, 0, None) / half_life_days)


# --------------------------------------------------------------------------
# Extração
# --------------------------------------------------------------------------
# Colunas que nem todo ano traz; DESCR_TIPOLOCAL falta em 2024 e é inferida do subtipo.
OPTIONAL_COLUMNS = {"DESCR_TIPOLOCAL"}


def fill_tipolocal(df: pd.DataFrame) -> pd.DataFrame:
    """Preenche DESCR_TIPOLOCAL ausente com o tipo mais comum do mesmo subtipo nos outros anos.

    246 dos 265 subtipos pertencem a um único tipo em >= 95% dos registros.
    """
    known = df.dropna(subset=["DESCR_TIPOLOCAL", "DESCR_SUBTIPOLOCAL"])
    if known.empty:
        return df
    modal = known.groupby("DESCR_SUBTIPOLOCAL")["DESCR_TIPOLOCAL"].agg(lambda s: s.value_counts().index[0])
    miss = df["DESCR_TIPOLOCAL"].isna()
    df.loc[miss, "DESCR_TIPOLOCAL"] = df.loc[miss, "DESCR_SUBTIPOLOCAL"].map(modal)
    print(f"DESCR_TIPOLOCAL inferido do subtipo em {int(miss.sum())} linhas "
          f"({int(df.loc[miss, 'DESCR_TIPOLOCAL'].isna().sum())} sem correspondência)")
    return df


def extract_city(xlsx_path) -> list[tuple]:
    """Lê as planilhas de dados (streaming) e devolve linhas de S.PAULO."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    out = []
    try:
        for name in wb.sheetnames:
            if name.upper().replace(" ", "_").startswith("CAMPOS_DA_TABELA"):
                continue
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            header = next(rows)
            # cabeçalhos variam entre anos (ex.: CIRCUNCRIÇÃO / CD_IBGE em 2025)
            idx = {norm_text(h): i for i, h in enumerate(header) if h}
            # coluna ausente no ano (ex.: DESCR_TIPOLOCAL em 2024) vira None e é inferida depois
            cols = [idx.get(norm_text(c)) for c in C.OCC_COLUMNS]
            missing = [c for c, i in zip(C.OCC_COLUMNS, cols, strict=True) if i is None]
            if missing and set(missing) - OPTIONAL_COLUMNS:
                raise KeyError(f"{xlsx_path.name}:{name}: colunas ausentes {missing}")
            i_city = next(i for k, i in idx.items() if k.startswith("nome_municipio_circun"))
            i_ibge = idx.get("cod ibge", idx.get("cd_ibge"))
            n = 0
            for r in rows:
                n += 1
                if r[i_city] == "S.PAULO" or r[i_ibge] == 3550308:
                    out.append(tuple(None if i is None else r[i] for i in cols))
            print(f"  {xlsx_path.name}:{name}: {n} linhas, acumulado SP {len(out)}")
    finally:
        wb.close()
    return out


def load_occurrences(force: bool = False) -> tuple[pd.DataFrame, list[str]]:
    """Ocorrências brutas da capital (cache em parquet)."""
    sources = sorted(C.DATA_RAW.glob("SPDadosCriminais_*.xlsx"))
    if not sources:
        raise SystemExit("Nenhuma planilha em data/raw. Rode pipeline.download_ssp antes.")
    newest = max(p.stat().st_mtime for p in sources)
    if not force and OCC_PATH.exists() and OCC_PATH.stat().st_mtime >= newest:
        print(f"Usando cache {OCC_PATH}")
        return pd.read_parquet(OCC_PATH), [p.name for p in sources]
    frames = []
    for p in sources:
        print(f"Lendo {p.name}")
        df = pd.DataFrame(extract_city(p), columns=C.OCC_COLUMNS)
        frames.append(df)
    df = pd.concat(frames, ignore_index=True)
    df = fill_tipolocal(df)
    # tipos heterogêneos (hora como time/str) -> normaliza antes do parquet
    df["HORA"] = df["HORA_OCORRENCIA_BO"].map(parse_hour).astype("Int8")
    df = df.drop(columns=["HORA_OCORRENCIA_BO"])
    df["DATA_OCORRENCIA_BO"] = pd.to_datetime(df["DATA_OCORRENCIA_BO"], errors="coerce")
    for c in ("LATITUDE", "LONGITUDE"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    for c in ("NUM_BO", "NOME_DELEGACIA", "DESC_PERIODO", "DESCR_TIPOLOCAL",
              "DESCR_SUBTIPOLOCAL", "RUBRICA", "DESCR_CONDUTA", "NATUREZA_APURADA", "BAIRRO"):
        df[c] = df[c].astype("string")
    df["ANO_BO"] = pd.to_numeric(df["ANO_BO"], errors="coerce").astype("Int32")
    C.DATA_PROCESSED.mkdir(parents=True, exist_ok=True)
    df.to_parquet(OCC_PATH, index=False)
    print(f"Cache gravado: {OCC_PATH} ({len(df)} linhas)")
    return df, [p.name for p in sources]


# --------------------------------------------------------------------------
# Modelo
# --------------------------------------------------------------------------
def prepare(df: pd.DataFrame, stats: dict | None = None) -> pd.DataFrame:
    """Classifica, filtra, deduplica e atribui período e peso temporal."""
    stats = stats if stats is not None else {}
    stats["raw_city"] = len(df)

    df = df.drop_duplicates(["NUM_BO", "ANO_BO", "NOME_DELEGACIA", "RUBRICA"]).copy()
    stats["after_dedupe"] = len(df)

    # classificação por combinações únicas (rápido)
    key_cols = ["RUBRICA", "DESCR_CONDUTA", "NATUREZA_APURADA", "DESCR_TIPOLOCAL", "DESCR_SUBTIPOLOCAL"]
    combos = df[key_cols].drop_duplicates()
    combos["category"] = [classify(*t) for t in combos.itertuples(index=False, name=None)]
    df = df.merge(combos, on=key_cols, how="left")
    df = df[df["category"].notna()]
    stats["after_classify"] = len(df)

    lat, lon = df["LATITUDE"], df["LONGITUDE"]
    min_lon, min_lat, max_lon, max_lat = C.SP_BBOX
    ok = (lat != 0) & (lon != 0) & lat.between(min_lat, max_lat) & lon.between(min_lon, max_lon)
    df = df[ok & df["DATA_OCORRENCIA_BO"].notna()].copy()
    stats["after_geo"] = len(df)

    d_max = df["DATA_OCORRENCIA_BO"].max()
    d_min = d_max - pd.Timedelta(days=C.WINDOW_DAYS)
    df = df[df["DATA_OCORRENCIA_BO"] > d_min].copy()
    stats["after_window"] = len(df)
    stats["date_min"] = df["DATA_OCORRENCIA_BO"].min().date().isoformat()
    stats["date_max"] = d_max.date().isoformat()

    age = (d_max - df["DATA_OCORRENCIA_BO"]).dt.days.to_numpy()
    df["w_t"] = decay_weight(age)
    hours = df["HORA"].to_numpy(dtype=float, na_value=np.nan)
    per = [period_from_hour(None if np.isnan(h) else int(h)) for h in hours]
    desc = df["DESC_PERIODO"].to_numpy()
    df["period"] = [p or period_from_desc(d) for p, d in zip(per, desc, strict=True)]
    return df.reset_index(drop=True)


def normalize_risk(score: np.ndarray) -> np.ndarray:
    """Escala logarítmica ancorada na mediana: risk = log1p(s/med) / log1p(p99.9/med).

    As contagens têm cauda longa (o Centro chega a ~40× a mediana); uma escala
    linear satura todo o Centro em 1.0 e o roteador deixa de distinguir ruas ali.
    """
    pos = score[score > 0]
    if pos.size == 0:
        return np.zeros_like(score, dtype=float)
    med = np.median(pos)
    top = np.percentile(pos, C.RISK_TOP_PERCENTILE)
    if top <= med:
        return np.minimum(1.0, score / top)
    return np.clip(np.log1p(score / med) / np.log1p(top / med), 0.0, 1.0)


def light_multiplier(light, gamma: float = C.RISK_GAMMA):
    """1 + γ·(1 - light); light NaN (sem dado) -> 1 (sem efeito)."""
    light = np.asarray(light, dtype=float)
    return np.where(np.isnan(light), 1.0, 1.0 + gamma * (1.0 - np.clip(light, 0.0, 1.0)))


def mix_risk(score, footfall, beta: float = C.RISK_BETA, f0: float | None = None):
    """Risco v2: mix = rate_norm^β · score_norm^(1-β), rate = score / (footfall + F0).

    F0 = mediana de footfall das células com score > 0 (se não informado). β=0 reproduz
    o v1 (normalize_risk(score)). Retorna (risk, rate).
    """
    score = np.asarray(score, dtype=float)
    ff = np.nan_to_num(np.asarray(footfall, dtype=float), nan=0.0)
    if f0 is None:
        f0 = float(np.median(ff[score > 0])) if (score > 0).any() else 0.0
    rate = score / (ff + max(f0, 1e-9))
    sn = normalize_risk(score)
    if beta <= 0:
        return sn, rate
    rn = normalize_risk(rate)
    return np.clip(rn ** beta * sn ** (1.0 - beta), 0.0, 1.0), rate


def aggregate(df: pd.DataFrame, footfall: pd.Series | None = None, light: pd.Series | None = None,
              beta: float = C.RISK_BETA, gamma: float = C.RISK_GAMMA) -> pd.DataFrame:
    """Agrega em H3 por (mode, period): suavização k=1, iluminação, taxa/mix e normalização log.

    `footfall`/`light`: Series indexadas por h3 (opcionais). Sem `light`, a coluna fica NaN.
    Em células sem ponto de luz (light ausente com base carregada) light = 0.
    """
    import h3

    cells = [h3.latlng_to_cell(a, o, C.H3_RES) for a, o in zip(df["LATITUDE"], df["LONGITUDE"], strict=True)]
    df = df.assign(h3=cells)
    out = []
    for mode in C.MODES:
        wts = C.WEIGHTS[mode]
        d = df.assign(wt=df["category"].map(wts) * df["w_t"])
        d = d[d["wt"] > 0]  # categoria com peso 0 no modo não contribui
        for period in C.PERIODS:
            sub = d if period == "all" else d[d["period"] == period]
            if sub.empty:
                continue
            g = sub.groupby("h3").agg(score=("wt", "sum"), n=("wt", "size"))
            cats = sub.groupby(["h3", "category"]).size()
            top = {h: json.dumps(s.droplevel(0).astype(int).to_dict(), ensure_ascii=False)
                   for h, s in cats.groupby(level=0)}
            base = g["score"].to_dict()
            score = dict(base)
            for cell, v in base.items():
                for nb in h3.grid_disk(cell, 1):
                    if nb != cell:
                        score[nb] = score.get(nb, 0.0) + C.NEIGHBOR_FACTOR * v
            res = pd.DataFrame({"h3": list(score), "score": list(score.values())})
            res["n"] = res["h3"].map(g["n"]).fillna(0).astype(int)
            res["top_categories"] = res["h3"].map(top).fillna("{}")
            if light is not None:
                res["light"] = res["h3"].map(light).fillna(0.0)
                if period in C.LIGHT_PERIODS:
                    res["score"] = res["score"] * light_multiplier(res["light"], gamma)
            else:
                res["light"] = np.nan
            if footfall is not None:
                res["footfall"] = res["h3"].map(footfall).fillna(0.0)
                res["risk"], res["rate"] = mix_risk(res["score"].to_numpy(), res["footfall"].to_numpy(), beta)
            else:
                res["footfall"], res["rate"] = np.nan, np.nan
                res["risk"] = normalize_risk(res["score"].to_numpy())
            res["mode"], res["period"] = mode, period
            out.append(res)
    r = pd.concat(out, ignore_index=True)
    return r[["h3", "mode", "period", "score", "risk", "n", "top_categories", "footfall", "light", "rate"]]


AUX_COLS = ["NUM_BO", "ANO_BO", "NOME_DELEGACIA", "NOME_MUNICIPIO", "RUBRICA", "LATITUDE", "LONGITUDE"]


def read_aux(path) -> pd.DataFrame:
    """Linhas de S.PAULO de uma base auxiliar (última planilha de dados), colunas AUX_COLS."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[-1]]
        rows = ws.iter_rows(values_only=True)
        header = list(next(rows))
        idx = [header.index(c) for c in AUX_COLS]
        df = pd.DataFrame([tuple(r[i] for i in idx) for r in rows], columns=AUX_COLS)
    finally:
        wb.close()
    return df[df["NOME_MUNICIPIO"] == "S.PAULO"].reset_index(drop=True)


def audit_aux_bases(main_df: pd.DataFrame) -> dict:
    """Mede a sobreposição das bases auxiliares da SSP com SPDadosCriminais.

    Chave do BO: (NUM_BO, ANO_BO, NOME_DELEGACIA). Reporta, por arquivo, quantos BOs
    (e linhas com coordenada válida) não existem em SPDadosCriminais.
    """
    main_df = main_df.dropna(subset=['NUM_BO', 'ANO_BO'])
    keys = set(zip(main_df["NUM_BO"].astype(str), main_df["ANO_BO"].astype(int),
                   main_df["NOME_DELEGACIA"].astype(str), strict=True))
    min_lon, min_lat, max_lon, max_lat = C.SP_BBOX
    report = {}
    for p in sorted(C.DATA_RAW.glob("*Subtraidos_*.xlsx")):
        a = read_aux(p)
        a["geo"] = pd.to_numeric(a["LATITUDE"], errors="coerce").between(min_lat, max_lat) & \
            pd.to_numeric(a["LONGITUDE"], errors="coerce").between(min_lon, max_lon)
        k = list(zip(a["NUM_BO"].astype(str), pd.to_numeric(a["ANO_BO"]).astype(int),
                     a["NOME_DELEGACIA"].astype(str), strict=True))
        a["in_main"] = [x in keys for x in k]
        bos = pd.Series(k).drop_duplicates()
        report[p.name] = {
            "rows_sp": int(len(a)), "bos_sp": int(len(bos)),
            "bos_in_main": int(sum(x in keys for x in bos)),
            "rows_not_in_main": int((~a["in_main"]).sum()),
            "rows_not_in_main_with_geo": int((~a["in_main"] & a["geo"]).sum()),
            # só o que o modelo usaria (roubo/furto com coordenada) e não está no principal
            "rows_not_in_main_roubo_furto_geo": int(
                (~a["in_main"] & a["geo"] & a["RUBRICA"].map(norm_text).str.contains(
                    "roubo \\(art|furto \\(art")).sum()),
        }
        print(p.name, report[p.name])
    return report


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--force", action="store_true", help="reextrai as planilhas")
    ap.add_argument("--audit-aux", action="store_true",
                    help="mede a sobreposição das bases Celulares/VeiculosSubtraidos e sai")
    ap.add_argument("--beta", type=float, default=C.RISK_BETA)
    ap.add_argument("--gamma", type=float, default=C.RISK_GAMMA)
    ap.add_argument("--no-footfall", action="store_true", help="v1: sem movimento (β efetivo 0)")
    ap.add_argument("--no-light", action="store_true")
    args = ap.parse_args()

    raw, sources = load_occurrences(args.force)
    if args.audit_aux:
        audit_aux_bases(raw)
        return
    stats: dict = {}
    df = prepare(raw, stats)
    from pipeline.footfall import load_footfall
    from pipeline.lighting import load_light

    footfall = None if args.no_footfall else load_footfall()
    light = None if args.no_light else load_light()
    risk = aggregate(df, footfall, light, args.beta, args.gamma)
    C.DATA_PROCESSED.mkdir(parents=True, exist_ok=True)
    risk.to_parquet(RISK_PATH, index=False)

    cat_counts = df["category"].value_counts().to_dict()
    meta = {
        "generated_at": dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
        "source_files": sources,
        "date_min": stats["date_min"],
        "date_max": stats["date_max"],
        "n_occurrences": int(len(df)),
        "categories": {k: int(v) for k, v in cat_counts.items()},
        "weights": C.WEIGHTS,
        "model": "v2" if footfall is not None else "v1",
        "beta": args.beta if footfall is not None else 0.0,
        "gamma": args.gamma if light is not None else 0.0,
        "footfall_source": "malha viária a pé do OSM ponderada por tipo de via (pipeline/footfall.py), H3 res 9, k=1" if footfall is not None else None,
        "light_source": ("GeoSampa WFS geoportal:iluminacao_publica (SMUL/ILUME), densidade H3 res 9 k=1"
                         if light is not None else None),
        "aux_bases": "Celulares/VeiculosSubtraidos não integradas: sem ocorrências novas (ver audit_aux_bases)",
    }
    META_PATH.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n== Resumo ==")
    for k in ("raw_city", "after_dedupe", "after_classify", "after_geo", "after_window"):
        print(f"{k:16s} {stats[k]:>9d}")
    print(f"Período: {stats['date_min']} a {stats['date_max']}")
    print("Categorias:", dict(Counter(cat_counts)))
    print("Células por modo/período:")
    print(risk.groupby(["mode", "period"]).size().unstack().to_string())
    print(f"Escrito: {RISK_PATH}, {META_PATH}")


if __name__ == "__main__":
    main()
